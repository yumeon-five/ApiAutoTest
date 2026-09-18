#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""gen_excel.py — 由中间 JSON（cases.json）生成接口测试用例汇总 Excel（.xlsx）。

规则（详见 reference/apicase-template-excel.md）：
- 每个接口单独一个 Sheet，依据 interface_key 分组区分。
- Sheet 名优先取 sheet_name → interface_name → interface_key；
  自动清洗为 Excel 合法（去除 []:*?/\\）、≤31 字符且不重复的名称。
- 缺失的 case_id 用 uuid.uuid4().hex（32 位无横杠）补全。
- 同时输出 接口_index.json：interfaces[key].caseIds，供 Markdown 回填用例编号。

依赖：openpyxl（若未安装：pip install openpyxl）

用法：
    python gen_excel.py <cases.json> [-o <output.xlsx>] [--index <接口_index.json>]
"""
import argparse
import json
import re
import sys
import uuid
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    sys.stderr.write("[gen_excel] 缺少依赖 openpyxl，请先执行：pip install openpyxl\n")
    sys.exit(1)

# 列顺序：(表头, cases.json 字段, 列宽)。与 reference/apicase-template-excel.md 一致。
COLUMNS = [
    ("接口Key", "interface_key", 16),
    ("用例编号", "case_id", 34),
    ("测试标题", "title", 30),
    ("前置条件", "precondition", 28),
    ("优先级", "priority", 8),
    ("请求URL", "url", 30),
    ("请求方法", "method", 10),
    ("请求头", "headers", 30),
    ("URL参数", "query_params", 24),
    ("请求体参数", "body_params", 36),
    ("预期结果", "expected", 46),
]

ILLEGAL_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
INDEX_FILENAME = "接口_index.json"


def to_cell(value):
    """把任意值转换为 openpyxl 可写入的单元格值。"""
    if value is None:
        return ""
    if isinstance(value, (str, int, float, bool)):
        return value
    # dict / list 等结构化为 JSON 字符串，避免写入报错
    return json.dumps(value, ensure_ascii=False)


def sanitize_sheet_name(name, used):
    """清洗 Sheet 名：去非法字符、截断 31 字符、保证唯一。"""
    cleaned = ILLEGAL_SHEET_CHARS.sub("_", str(name or "").strip())
    cleaned = cleaned[:31] or "Sheet"
    base = cleaned
    idx = 1
    while cleaned in used:
        suffix = "_{}".format(idx)
        cleaned = base[: 31 - len(suffix)] + suffix
        idx += 1
    used.add(cleaned)
    return cleaned


def build_workbook(data):
    """根据 cases.json 数据构建 Workbook，返回 (wb, index)。"""
    interfaces = data.get("interfaces")
    if not isinstance(interfaces, dict) or not interfaces:
        raise ValueError("cases.json 缺少非空的 interfaces 字典（须按 interface_key 分组）")

    wb = Workbook()
    wb.remove(wb.active)  # 删除默认 Sheet，改为每接口一个

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4F81BD")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    body_align = Alignment(wrap_text=True, vertical="top")

    used_names = set()
    index = {"interfaces": {}}

    for key, iface in interfaces.items():
        iface = iface or {}
        cases = iface.get("cases") or []
        sheet_label = iface.get("sheet_name") or iface.get("interface_name") or key
        ws = wb.create_sheet(sanitize_sheet_name(sheet_label, used_names))

        # 表头
        for col_idx, (title, _field, width) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=1, column=col_idx, value=title)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            ws.column_dimensions[get_column_letter(col_idx)].width = width
        ws.freeze_panes = "A2"

        # 用例行
        case_ids = []
        for row_idx, case in enumerate(cases, start=2):
            case = case or {}
            case_id = str(case.get("case_id") or "").strip()
            if not case_id:
                case_id = uuid.uuid4().hex  # 规范要求：32 位无横杠
            case_ids.append(case_id)

            for col_idx, (_title, field, _width) in enumerate(COLUMNS, start=1):
                value = key if field == "interface_key" else case.get(field, "")
                if field == "case_id":
                    value = case_id
                ws.cell(row=row_idx, column=col_idx, value=to_cell(value)).alignment = body_align

        index["interfaces"][key] = {
            "interface_name": iface.get("interface_name", ""),
            "caseIds": case_ids,
        }

    if not wb.sheetnames:
        raise ValueError("没有生成任何接口 Sheet，请检查 cases.json 的 interfaces 是否为空")
    return wb, index


def resolve_out_path(args, data, src):
    """确定输出 xlsx 路径。"""
    out = args.output or data.get("output_xlsx")
    if out:
        out_path = Path(out)
        if not out_path.is_absolute():
            out_path = src.parent / out_path
    else:
        out_path = src.with_suffix(".xlsx")
    if out_path.suffix.lower() != ".xlsx":
        out_path = out_path.with_suffix(".xlsx")
    return out_path


def main(argv=None):
    parser = argparse.ArgumentParser(description="由中间 JSON 生成接口测试用例汇总 Excel")
    parser.add_argument("cases_json", help="中间 JSON 文件路径（cases.json）")
    parser.add_argument("-o", "--output", help="输出 .xlsx 路径；缺省取 JSON 的 output_xlsx 或同名 .xlsx")
    parser.add_argument("--index", help="接口索引输出路径；缺省为 xlsx 同目录下 接口_index.json")
    args = parser.parse_args(argv)

    src = Path(args.cases_json)
    if not src.exists():
        sys.stderr.write("[gen_excel] 找不到输入文件：{}\n".format(src))
        return 1

    try:
        data = json.loads(src.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        sys.stderr.write("[gen_excel] cases.json 解析失败：{}\n".format(exc))
        return 1

    try:
        wb, index = build_workbook(data)
    except ValueError as exc:
        sys.stderr.write("[gen_excel] {}\n".format(exc))
        return 1

    out_path = resolve_out_path(args, data, src)
    if args.index:
        # 仅对用户显式传入的相对 --index 以 cases.json 所在目录重新定位；
        # 默认路径已由 out_path.parent 正确定位，不能再拼接，否则目录会重复一层。
        index_path = Path(args.index)
        if not index_path.is_absolute():
            index_path = src.parent / index_path
    else:
        index_path = out_path.parent / INDEX_FILENAME

    try:
        out_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(out_path)
        index_path.parent.mkdir(parents=True, exist_ok=True)
        index_path.write_text(
            json.dumps(index, ensure_ascii=False, indent=2), encoding="utf-8"
        )
    except PermissionError:
        sys.stderr.write("[gen_excel] 写入被拒绝，请先关闭已打开的 Excel 文件：{}\n".format(out_path))
        return 1

    total = sum(len(v["caseIds"]) for v in index["interfaces"].values())
    print("[gen_excel] 已生成 Excel：{}".format(out_path))
    print("[gen_excel] Sheet 数（接口数）：{}，用例总数：{}".format(len(index["interfaces"]), total))
    print("[gen_excel] 已生成接口索引：{}".format(index_path))
    return 0


if __name__ == "__main__":
    sys.exit(main())
